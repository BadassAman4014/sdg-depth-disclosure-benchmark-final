#!/usr/bin/env python3
"""
create_empty_human_coding_sheets.py — Prepares completely empty 50-sample coding sheets for human labeling.
"""

import sys
import pandas as pd
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
SAMPLES_PATH = DATA_DIR / "samples_50.xlsx"


def create_and_style_sheet(df_samples: pd.DataFrame, file_path: Path, sheet_type: str):
    df = df_samples.copy()
    
    if sheet_type == "single":
        # Master single coder sheet
        df["human_score_0_to_5"] = ""
        df["has_target (Yes/No)"] = ""
        df["has_action (Yes/No)"] = ""
        df["has_measured_outcome (Yes/No)"] = ""
        df["notes_and_evidence"] = ""
        header_color = "1E3A8A"  # Deep blue
    elif sheet_type == "coder_a":
        df["person_A_score_0_to_5"] = ""
        df["has_target (Yes/No)"] = ""
        df["has_action (Yes/No)"] = ""
        df["has_measured_outcome (Yes/No)"] = ""
        df["notes_and_evidence"] = ""
        header_color = "065F46"  # Forest green
    elif sheet_type == "coder_b":
        df["person_B_score_0_to_5"] = ""
        df["has_target (Yes/No)"] = ""
        df["has_action (Yes/No)"] = ""
        df["has_measured_outcome (Yes/No)"] = ""
        df["notes_and_evidence"] = ""
        header_color = "7C2D12"  # Deep amber/burgundy

    df.to_excel(file_path, index=False)

    # Style with OpenPyXL
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    # Input column highlight (soft yellow) for empty human score column
    input_fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")

    score_col_idx = None
    for col in range(1, ws.max_column + 1):
        h_val = str(ws.cell(row=1, column=col).value).lower()
        if "score" in h_val:
            score_col_idx = col

    # Header styling
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws.row_dimensions[1].height = 30

    # Data row styling
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 70
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = data_font
            cell.border = cell_border
            header_val = str(ws.cell(row=1, column=col).value).lower()
            
            if "passage" in header_val:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            elif any(k in header_val for k in ["score", "id", "year", "target", "action", "outcome"]):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Highlight the blank score cell so human knows where to type
            if col == score_col_idx:
                cell.fill = input_fill
                cell.font = Font(name="Segoe UI", size=12, bold=True, color="1E3A8A")

    # Column widths
    for col in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col)
        h = str(ws.cell(row=1, column=col).value).lower()
        if "passage" in h:
            ws.column_dimensions[col_letter].width = 65
        elif "notes" in h:
            ws.column_dimensions[col_letter].width = 30
        elif "score" in h:
            ws.column_dimensions[col_letter].width = 22
        elif "id" in h or "year" in h:
            ws.column_dimensions[col_letter].width = 12
        elif "keyword" in h or "company" in h:
            ws.column_dimensions[col_letter].width = 22
        else:
            ws.column_dimensions[col_letter].width = 18

    wb.save(file_path)
    print(f"Created empty coding sheet: {file_path}")


def main():
    if not SAMPLES_PATH.exists():
        print(f"Error: {SAMPLES_PATH} not found.")
        sys.exit(1)

    df_samples = pd.read_excel(SAMPLES_PATH)

    # 1. Master Single Human Coding Sheet (For you to label)
    create_and_style_sheet(df_samples, DATA_DIR / "coding_sheet_50_samples.xlsx", "single")

    # 2. Coder A Sheet (For Person A)
    create_and_style_sheet(df_samples, DATA_DIR / "coding_sheet_person_A_50.xlsx", "coder_a")

    # 3. Coder B Sheet (For Person B)
    create_and_style_sheet(df_samples, DATA_DIR / "coding_sheet_person_B_50.xlsx", "coder_b")

    print("\n✅ All human coding sheets are prepared with completely blank label columns ready for human evaluation!")


if __name__ == "__main__":
    main()
