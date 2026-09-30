#!/usr/bin/env python3
"""
run_pipeline_50.py — End-to-End Pipeline Execution on 50 Samples (0 to 5 Depth Scoring).
Reference: Izhar et al. (2026) Table 1 / Hummel (2019) / PwC (2018).

Usage:
  python run_pipeline_50.py
"""

import sys
import re
import joblib
import numpy as np
import pandas as pd
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Import continuous pipeline module
from continuous_model import ContinuousDepthPipeline, extract_linguistic_features, train_and_save_model

sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = Path(__file__).resolve().parent
INPUT_SAMPLES = BASE_DIR / "data" / "samples_50.xlsx"
GOLDEN_TRUTH = BASE_DIR / "data" / "golden_truth_50.xlsx"
OUTPUT_EXCEL = BASE_DIR / "output_predictions_50.xlsx"
OUTPUT_CSV = BASE_DIR / "output_predictions_50.csv"
MODEL_PATH = BASE_DIR / "model_depth_50.joblib"

# Regex rules for Izhar et al. (Table 1) linguistic components
QUAL_TARGET_RE = re.compile(r"\b(?:strive|aim|aspire|commit|intend|pledge|endeavor|target|vision|seek|goal|plan|streben|beabsichtigen)\b", re.I)
QUANT_TARGET_RE = re.compile(r"\b(?:target|aim|reduce|cut|reach)\s+(?:by|to|of)?\s*\d+(?:[\.,]\d+)?\s*(?:%|percent|prozent|tco2|kwh|tons?|tonnes?|eur|€|\$|by\s+20\d\d)\b", re.I)
ACTION_RE = re.compile(r"\b(?:installed|implemented|deployed|commissioned|invested|allocated|constructed|built|initiated|trained|upgraded|equipped|umgesetzt|eingeführt|investiert|reduziert)\b", re.I)
QUAL_OUTCOME_RE = re.compile(r"\b(?:resulted in|led to|improved|enhanced|strengthened|achieved progress|positive impact|verbessert|gestärkt|erfolgreich)\b", re.I)
QUANT_OUTCOME_RE = re.compile(r"\b(?:reduced|decreased|cut|saved|generated|sourced)\s+(?:by\s+)?\d+(?:[\.,]\d+)?\s*(?:%|percent|prozent|tco2|tco2e|kwh|mwh|gwh|tonnes?|million|billion|mio|mrd|€|\$)\b", re.I)
AUDIT_RE = re.compile(r"\b(?:verified|assured|audited|certified)\s+by\b|\b(?:pwc|kpmg|ey|deloitte|tüv|tuev|bureau\s+veritas|iso\s*14001|iso\s*50001)\b", re.I)


def extract_components(text: str) -> dict:
    p_str = str(text)
    return {
        "has_qual_target": bool(QUAL_TARGET_RE.search(p_str)),
        "has_quant_target": bool(QUANT_TARGET_RE.search(p_str)),
        "has_action": bool(ACTION_RE.search(p_str)),
        "has_qual_outcome": bool(QUAL_OUTCOME_RE.search(p_str)),
        "has_quant_outcome": bool(QUANT_OUTCOME_RE.search(p_str)),
        "has_audit": bool(AUDIT_RE.search(p_str)),
    }


def generate_izhar_rationale(score_int: int, comp: dict) -> str:
    """Provides explicit theoretical justification based on Izhar et al. Table 1."""
    if score_int == 5:
        return "Score 5: Quantitative target and/or third-party audit with verified outcome measurement."
    elif score_int == 4:
        return "Score 4: Qualitative target with reported quantitative measurement of outcome."
    elif score_int == 3:
        return "Score 3: Qualitative target with qualitative measurement of outcome."
    elif score_int == 2:
        if comp["has_quant_target"]:
            return "Score 2: Standalone quantitative target reported."
        else:
            return "Score 2: Qualitative target accompanied by concrete implemented SDG actions."
    elif score_int == 1:
        return "Score 1: Aspirational qualitative target or policy statement without operational outcome."
    else:
        return "Score 0: Pure boilerplate / general disclosure without target, action, or measured outcome."


def style_output_excel(file_path: Path):
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    # Color highlights for scores
    score_fills = {
        0: PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid"),  # Light gray
        1: PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),  # Light yellow
        2: PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid"),  # Light blue
        3: PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),  # Soft green
        4: PatternFill(start_color="BBF7D0", end_color="BBF7D0", fill_type="solid"),  # Vibrant green
        5: PatternFill(start_color="86EFAC", end_color="86EFAC", fill_type="solid"),  # Deep green
    }

    # Style Header
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
    ws.row_dimensions[1].height = 28

    # Style Data
    score_col_idx = None
    for col in range(1, ws.max_column + 1):
        if "izhar_level" in str(ws.cell(row=1, column=col).value).lower():
            score_col_idx = col
            break

    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 65
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = data_font
            cell.border = cell_border
            header_val = str(ws.cell(row=1, column=col).value).lower()
            if "passage" in header_val or "rationale" in header_val:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            elif any(k in header_val for k in ["score", "id", "year", "target", "action", "outcome", "audit", "label", "level"]):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        # Highlight score column
        if score_col_idx:
            val = ws.cell(row=row, column=score_col_idx).value
            try:
                int_val = int(val)
                if int_val in score_fills:
                    ws.cell(row=row, column=score_col_idx).fill = score_fills[int_val]
                    ws.cell(row=row, column=score_col_idx).font = Font(name="Segoe UI", size=11, bold=True)
            except (ValueError, TypeError):
                pass

    # Column widths
    for col in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col)
        h = str(ws.cell(row=1, column=col).value).lower()
        if "passage" in h:
            ws.column_dimensions[col_letter].width = 65
        elif "rationale" in h:
            ws.column_dimensions[col_letter].width = 45
        elif "continuous" in h:
            ws.column_dimensions[col_letter].width = 18
        elif "level" in h or "score" in h or "id" in h:
            ws.column_dimensions[col_letter].width = 14
        elif "keyword" in h or "company" in h:
            ws.column_dimensions[col_letter].width = 22
        else:
            ws.column_dimensions[col_letter].width = 15

    wb.save(file_path)


def main():
    print("=" * 85)
    print("  SDG DEPTH PIPELINE TEST — 50 SAMPLES (Izhar et al. 2026 Table 1 Scoring)")
    print("=" * 85)

    if not INPUT_SAMPLES.exists():
        print(f"Error: Input samples not found at {INPUT_SAMPLES}. Please run setup_test_data.py first.")
        sys.exit(1)

    print(f"1. Loading 50 samples from: {INPUT_SAMPLES.name}...")
    df = pd.read_excel(INPUT_SAMPLES)
    passages = df["passage"].fillna("").astype(str).tolist()

    if not MODEL_PATH.exists():
        print(f"Model artifact not found at {MODEL_PATH.name}. Training directly from golden truth...")
        train_and_save_model(GOLDEN_TRUTH, MODEL_PATH)

    print(f"2. Loading trained continuous regressor ({MODEL_PATH.name})...")
    pipeline = joblib.load(MODEL_PATH)

    print("3. Predicting continuous depth scores (0.00 to 5.00)...")
    continuous_scores = pipeline.predict(passages)

    print("4. Extracting Izhar et al. (Table 1) linguistic components & rationales...")
    comp_records = []
    rationales = []
    integer_levels = []
    binary_labels = []

    for p, cont_s in zip(passages, continuous_scores):
        comp = extract_components(p)
        comp_records.append(comp)

        # Discretize continuous score into Table 1 integer level (0 to 5)
        int_s = int(np.clip(np.round(cont_s), 0, 5))
        integer_levels.append(int_s)
        binary_labels.append("sub" if int_s >= 2 else "sym")
        rationales.append(generate_izhar_rationale(int_s, comp))

    comp_df = pd.DataFrame(comp_records)

    # Build Output DataFrame
    results_df = pd.DataFrame({
        "sample_id": df["sample_id"],
        "company": df["company"],
        "year": df["year"],
        "sdg_category": df["sdg_category"],
        "keyword": df["keyword"],
        "passage": df["passage"],
        "continuous_score (0.00-5.00)": np.round(continuous_scores, 2),
        "izhar_level (0-5)": integer_levels,
        "binary_label": binary_labels,
        "has_target": (comp_df["has_qual_target"] | comp_df["has_quant_target"]).map({True: "Yes", False: "No"}),
        "has_action": comp_df["has_action"].map({True: "Yes", False: "No"}),
        "has_measured_outcome": (comp_df["has_qual_outcome"] | comp_df["has_quant_outcome"]).map({True: "Yes", False: "No"}),
        "has_audit_assurance": comp_df["has_audit"].map({True: "Yes", False: "No"}),
        "scoring_rationale": rationales,
    })

    # Save outputs
    results_df.to_excel(OUTPUT_EXCEL, index=False)
    style_output_excel(OUTPUT_EXCEL)
    results_df.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")

    print(f"\n✅ Pipeline executed successfully!")
    print(f"  • Excel Output: {OUTPUT_EXCEL}")
    print(f"  • CSV Output:   {OUTPUT_CSV}")

    # Summary Statistics
    print("\n" + "=" * 85)
    print("  SCORE DISTRIBUTION ACROSS 50 SAMPLES (Izhar et al. Table 1)")
    print("=" * 85)
    level_counts = pd.Series(integer_levels).value_counts().sort_index()
    for lvl in range(6):
        cnt = level_counts.get(lvl, 0)
        pct = (cnt / len(df)) * 100
        desc = [
            "No SDG Information / Boilerplate",
            "Qualitative Target Only (Aspiration)",
            "Qual Target + Action OR Quant Target",
            "Qual Target + Qual Outcome",
            "Qual Target + Quant Outcome",
            "Quant Target + Measured/Audited Outcome",
        ][lvl]
        bar = "█" * int(pct // 3)
        print(f"  Score {lvl} | {cnt:2d} passages ({pct:5.1f}%) | {bar:<15} | {desc}")

    print("-" * 85)
    sym_cnt = (results_df["binary_label"] == "sym").sum()
    sub_cnt = (results_df["binary_label"] == "sub").sum()
    print(f"  Overall Binary Mapping:  Symbolic (Scores 0-1): {sym_cnt} ({sym_cnt*2}%) | Substantive (Scores 2-5): {sub_cnt} ({sub_cnt*2}%)")
    print(f"  Average SDG Depth Score: {np.mean(continuous_scores):.2f} / 5.00")
    print("=" * 85)

    print("\n  Sample First 5 Predictions:")
    print("  " + "-" * 80)
    for i in range(5):
        row = results_df.iloc[i]
        p_short = str(row["passage"])[:70].replace("\n", " ")
        print(f"  Row {row['sample_id']:2d} | Score: {row['continuous_score (0.00-5.00)']:.2f} (Level {row['izhar_level (0-5)']}) [{row['binary_label'].upper()}] -> \"{p_short}...\"")
    print("=" * 85 + "\n")


if __name__ == "__main__":
    main()
